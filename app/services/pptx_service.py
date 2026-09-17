from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.services.analytics_service import AnalyticsService


class PPTXService:

    def generate_sales_analysis_deck(
        self,
        start_date,
        end_date,
    ) -> dict:
        """
        Generate a PowerPoint sales analysis deck
        for the given date range.
        """

        try:
            # -----------------------------------------------------
            # 1. Get analytics data
            # -----------------------------------------------------
            analytics_service = AnalyticsService()

            result = analytics_service.get_sales_analysis_data(
                start_date=start_date,
                end_date=end_date,
            )

            if not result.get("success"):
                return result

            data = result["data"]

            summary = data["summary"]
            daily_sales = data["daily_sales"]
            payment_methods = data["payment_methods"]
            top_products = data["top_products"]

            # -----------------------------------------------------
            # 2. Create output directory
            # -----------------------------------------------------
            output_directory = Path(
                "generated",
                "reports",
            )

            output_directory.mkdir(
                parents=True,
                exist_ok=True,
            )

            output_path = (
                output_directory
                / (
                    f"sales_analysis_"
                    f"{start_date.isoformat()}_"
                    f"{end_date.isoformat()}.pptx"
                )
            )

            # -----------------------------------------------------
            # 3. Create presentation
            # -----------------------------------------------------
            presentation = Presentation()

            # -----------------------------------------------------
            # Slide 1 — Title
            # -----------------------------------------------------
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[0]
            )

            slide.shapes.title.text = "KiranaAI Sales Analysis"

            subtitle = slide.placeholders[1]

            subtitle.text = (
                f"{start_date.strftime('%d-%m-%Y')} "
                f"to "
                f"{end_date.strftime('%d-%m-%Y')}"
            )

            # -----------------------------------------------------
            # Slide 2 — Executive Summary
            # -----------------------------------------------------
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )

            slide.shapes.title.text = "Executive Summary"

            textbox = slide.shapes.add_textbox(
                Inches(1),
                Inches(1.5),
                Inches(8),
                Inches(4.5),
            )

            text_frame = textbox.text_frame
            text_frame.clear()

            summary_lines = [
                f"Total Sales: ₹{summary['total_sales']:.2f}",
                f"Total Bills: {summary['bill_count']}",
                f"Subtotal: ₹{summary['subtotal']:.2f}",
                f"CGST: ₹{summary['cgst']:.2f}",
                f"SGST: ₹{summary['sgst']:.2f}",
                f"Tax Collected: ₹{summary['tax_collected']:.2f}",
            ]

            for index, line in enumerate(summary_lines):

                if index == 0:
                    paragraph = text_frame.paragraphs[0]
                else:
                    paragraph = text_frame.add_paragraph()

                paragraph.text = line
                paragraph.font.size = Pt(22)
                paragraph.space_after = Pt(12)

            # -----------------------------------------------------
            # Slide 3 — Daily Sales Trend
            # -----------------------------------------------------
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )

            slide.shapes.title.text = "Daily Sales Trend"

            if daily_sales:

                chart_data = ChartData()

                chart_data.categories = [
                    row["date"]
                    for row in daily_sales
                ]

                chart_data.add_series(
                    "Sales",
                    [
                        row["sales"]
                        for row in daily_sales
                    ],
                )

                chart = slide.shapes.add_chart(
                    XL_CHART_TYPE.LINE,
                    Inches(0.8),
                    Inches(1.5),
                    Inches(8.5),
                    Inches(4.8),
                    chart_data,
                ).chart

                chart.has_legend = False
                chart.has_title = False

                chart.value_axis.has_major_gridlines = True

            else:

                textbox = slide.shapes.add_textbox(
                    Inches(1),
                    Inches(2),
                    Inches(8),
                    Inches(2),
                )

                textbox.text_frame.text = (
                    "No finalized sales found for this period."
                )

            # -----------------------------------------------------
            # Slide 4 — Payment Method Breakdown
            # -----------------------------------------------------
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )

            slide.shapes.title.text = (
                "Payment Method Breakdown"
            )

            available_payments = {
                method: amount
                for method, amount in payment_methods.items()
                if amount > 0
            }

            if available_payments:

                chart_data = ChartData()

                chart_data.categories = list(
                    available_payments.keys()
                )

                chart_data.add_series(
                    "Amount",
                    list(
                        available_payments.values()
                    ),
                )

                chart = slide.shapes.add_chart(
                    XL_CHART_TYPE.PIE,
                    Inches(1.5),
                    Inches(1.4),
                    Inches(7),
                    Inches(5),
                    chart_data,
                ).chart

                chart.has_legend = True
                chart.has_title = False

            else:

                textbox = slide.shapes.add_textbox(
                    Inches(1),
                    Inches(2),
                    Inches(8),
                    Inches(2),
                )

                textbox.text_frame.text = (
                    "No finalized payments found for this period."
                )

            # -----------------------------------------------------
            # Slide 5 — Top-Selling Products
            # -----------------------------------------------------
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )

            slide.shapes.title.text = (
                "Top-Selling Products"
            )

            if top_products:

                chart_data = ChartData()

                chart_data.categories = [
                    row["product_name"]
                    for row in top_products[:5]
                ]

                chart_data.add_series(
                    "Quantity Sold",
                    [
                        row["quantity_sold"]
                        for row in top_products[:5]
                    ],
                )

                chart = slide.shapes.add_chart(
                    XL_CHART_TYPE.BAR_CLUSTERED,
                    Inches(1),
                    Inches(1.4),
                    Inches(8),
                    Inches(5),
                    chart_data,
                ).chart

                chart.has_legend = False
                chart.has_title = False

            else:

                textbox = slide.shapes.add_textbox(
                    Inches(1),
                    Inches(2),
                    Inches(8),
                    Inches(2),
                )

                textbox.text_frame.text = (
                    "No product sales found for this period."
                )

            # -----------------------------------------------------
            # Slide 6 — Product Sales Details
            # -----------------------------------------------------
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )

            slide.shapes.title.text = (
                "Product Sales Details"
            )

            if top_products:

                rows = min(
                    len(top_products),
                    6,
                ) + 1

                columns = 4

                table = slide.shapes.add_table(
                    rows,
                    columns,
                    Inches(0.7),
                    Inches(1.5),
                    Inches(8.8),
                    Inches(4.5),
                ).table

                headers = [
                    "Product",
                    "Unit",
                    "Qty Sold",
                    "Sales",
                ]

                for column, header in enumerate(headers):

                    cell = table.cell(
                        0,
                        column,
                    )

                    cell.text = header

                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.bold = True
                        paragraph.font.size = Pt(12)
                        paragraph.alignment = (
                            PP_ALIGN.CENTER
                        )

                for row_index, product in enumerate(
                    top_products[:6],
                    start=1,
                ):

                    values = [
                        product["product_name"],
                        product["unit"],
                        f"{product['quantity_sold']:.2f}",
                        f"₹{product['sales_amount']:.2f}",
                    ]

                    for column, value in enumerate(values):

                        cell = table.cell(
                            row_index,
                            column,
                        )

                        cell.text = value

                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.size = Pt(11)

            else:

                textbox = slide.shapes.add_textbox(
                    Inches(1),
                    Inches(2),
                    Inches(8),
                    Inches(2),
                )

                textbox.text_frame.text = (
                    "No product sales found for this period."
                )

            # -----------------------------------------------------
            # Slide 7 — Closing Insights
            # -----------------------------------------------------
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )

            slide.shapes.title.text = "Business Summary"

            textbox = slide.shapes.add_textbox(
                Inches(1),
                Inches(1.5),
                Inches(8),
                Inches(4.5),
            )

            text_frame = textbox.text_frame
            text_frame.clear()

            insights = []

            if summary["total_sales"] > 0:
                insights.append(
                    f"Sales during the period were "
                    f"₹{summary['total_sales']:.2f}."
                )

            if summary["bill_count"] > 0:
                average_bill = (
                    summary["total_sales"]
                    / summary["bill_count"]
                )

                insights.append(
                    f"Average bill value was "
                    f"₹{average_bill:.2f}."
                )

            if top_products:

                top_product = top_products[0]

                insights.append(
                    f"Top-selling product by quantity was "
                    f"{top_product['product_name']} "
                    f"with "
                    f"{top_product['quantity_sold']:.2f} "
                    f"{top_product['unit']} sold."
                )

            if payment_methods:

                largest_payment = max(
                    payment_methods,
                    key=payment_methods.get,
                )

                if payment_methods[largest_payment] > 0:

                    insights.append(
                        f"{largest_payment} accounted for "
                        f"₹{payment_methods[largest_payment]:.2f} "
                        f"in finalized sales."
                    )

            if not insights:
                insights.append(
                    "No finalized sales were recorded "
                    "during this period."
                )

            for index, insight in enumerate(insights):

                if index == 0:
                    paragraph = text_frame.paragraphs[0]
                else:
                    paragraph = text_frame.add_paragraph()

                paragraph.text = f"• {insight}"
                paragraph.font.size = Pt(20)
                paragraph.space_after = Pt(15)

            # -----------------------------------------------------
            # 8. Save presentation
            # -----------------------------------------------------
            presentation.save(str(output_path))

            return {
                "success": True,
                "file_path": str(output_path),
                "start_date": start_date.isoformat(),
                "end_date": end_date.isoformat(),
                "total_sales": summary["total_sales"],
                "bill_count": summary["bill_count"],
            }

        except Exception as error:

            return {
                "success": False,
                "error": str(error),
            }